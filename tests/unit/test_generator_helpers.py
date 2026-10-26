"""
Pure helpers in generator.py: competency grouping,
uploaded-file text extraction and prompt fragments.
"""

from collections import Counter

import pytest

from assessment.batching import Batch
from assessment.generator import (
    _competency_bucket,
    _format_batch_time_limit,
    _format_final_batch_total,
    compute_blooms_by_type,
    extract_docx_text_sync,
    extract_uploaded_text,
    format_previously_generated,
    format_question_type_instructions,
)


# --------------------------------------------------------------------------
# compute_blooms_by_type / prompt fragments
# --------------------------------------------------------------------------

def test_compute_blooms_by_type_assigns_one_level_per_requested_question():
    out = compute_blooms_by_type({"Remember": 50, "Apply": 50}, {"mcq": 3, "ftb": 1, "mtf": 0})
    assert set(out) == {"mcq", "ftb"}
    assert len(out["mcq"]) == 3 and len(out["ftb"]) == 1
    assert Counter(l for levels in out.values() for l in levels) == {"Remember": 2, "Apply": 2}


def test_question_type_instructions_mark_unrequested_types_do_not_generate():
    text = format_question_type_instructions({"mcq": 2})
    assert "2 Multiple Choice Questions" in text
    assert text.count("[DO NOT GENERATE]") == 4


def test_previously_generated_lists_stems_grouped_by_type():
    digest = [
        {"type": "ftb", "stem": "F stem", "blooms_level": "Apply", "course_name": "C1"},
        {"type": "mcq", "stem": "M stem", "blooms_level": "", "course_name": ""},
    ]
    text = format_previously_generated(digest)
    assert "DO NOT REPEAT" in text
    assert text.index("M stem") < text.index("F stem")  # MCQ group first
    assert "[Apply · C1] F stem" in text
    assert format_previously_generated([]) == ""


def test_batch_time_limit_states_whole_assessment_and_per_question_pace():
    text = _format_batch_time_limit("120 minutes", 200)
    assert "COMPLETE assessment of 200" in text
    assert "0.6 minute(s) per question" in text
    assert _format_batch_time_limit(None, 10) is None


def test_final_batch_total_separates_batch_counts_from_blueprint_totals():
    batch = Batch(index=1, type_counts={"mcq": 5}, is_final=True)
    text = _format_final_batch_total(batch, 30, {"mcq": 20, "ftb": 10, "mtf": 0})
    assert text.startswith("5 for THIS PART")
    assert "whole assessment is 30" in text
    assert "MTF" not in text


# --------------------------------------------------------------------------
# _competency_bucket (commit a057f1c)
# --------------------------------------------------------------------------

@pytest.mark.parametrize("area,expected", [
    ("Behavioural", "behavioral"),
    ("behavioral", "behavioral"),
    ("Behavioural - Core", "behavioral"),
    ("  Functional ", "functional"),
    ("Domain", "domain"),
    ("", "domain"),
    (None, "domain"),
])
def test_competency_bucket(area, expected):
    assert _competency_bucket(area) == expected


# --------------------------------------------------------------------------
# Uploaded file extraction (commit a4f5a7a)
# --------------------------------------------------------------------------

async def test_txt_upload_strips_bom_and_nul(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_bytes("﻿Hello\x00 world\n".encode("utf-8"))
    assert await extract_uploaded_text(path) == "Hello world"


async def test_txt_upload_falls_back_to_latin1(tmp_path):
    path = tmp_path / "legacy.txt"
    path.write_bytes("caf\xe9".encode("latin-1"))
    assert await extract_uploaded_text(path) == "café"


async def test_extension_match_is_case_insensitive(tmp_path):
    path = tmp_path / "NOTES.TXT"
    path.write_text("upper", encoding="utf-8")
    assert await extract_uploaded_text(path) == "upper"


async def test_vtt_upload_drops_cues_and_timestamps(tmp_path):
    path = tmp_path / "t.vtt"
    path.write_text("WEBVTT\n\n1\n00:00:01.000 --> 00:00:02.000\nHello there\n// note\n", encoding="utf-8")
    assert await extract_uploaded_text(path) == "Hello there"


async def test_docx_upload_reads_paragraphs_and_tables(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("   ")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Cell A"
    table.rows[0].cells[1].text = "Cell B"
    path = tmp_path / "doc.docx"
    document.save(str(path))

    assert await extract_uploaded_text(path) == "First paragraph\nCell A | Cell B"


def test_corrupt_docx_returns_empty_string_instead_of_raising(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")
    assert extract_docx_text_sync(path) == ""


async def test_pdf_upload_extracts_text(tmp_path):
    import fitz

    path = tmp_path / "doc.pdf"
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Hello PDF")
    pdf.save(str(path))
    pdf.close()
    assert "Hello PDF" in await extract_uploaded_text(path)


async def test_unsupported_upload_type_is_skipped(tmp_path):
    path = tmp_path / "sheet.xlsx"
    path.write_bytes(b"whatever")
    assert await extract_uploaded_text(path) == ""
